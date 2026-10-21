"""Flask-интерфейс SEO Meta Auditor."""
import hmac
import io
import json
import os
import re
import secrets
import time

import requests
from flask import Flask, Response, render_template, request, send_file, session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import auditor

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
RESULTS_DIR = os.path.join(BASE_DIR, "results")
RESULTS_TTL = 24 * 3600  # результаты старше суток удаляются

try:  # необязательный локальный файл с логином/паролем (в git не попадает)
    import config_local
except ImportError:
    config_local = None


def setting(name, default=""):
    """Значение из переменной окружения или из config_local.py."""
    return os.environ.get(name) or getattr(config_local, name, default)


def load_secret_key():
    key = setting("SECRET_KEY")
    if key:
        return key
    path = os.path.join(BASE_DIR, ".secret_key")
    if not os.path.exists(path):
        with open(path, "w") as f:
            f.write(secrets.token_hex(32))
    with open(path) as f:
        return f.read().strip()


app = Flask(__name__)
app.secret_key = load_secret_key()
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax")
app.config["AUTH_REQUIRED"] = True  # при локальном запуске без пароля отключается ниже

FILLS = {"critical": "F8D7DA", "warning": "FFF3CD", "info": "E2E3E5", "ok": "D4EDDA"}


@app.before_request
def require_auth():
    if not app.config["AUTH_REQUIRED"]:
        return None
    user, password = setting("APP_USER"), setting("APP_PASSWORD")
    if not password:
        return Response("Не задан APP_PASSWORD: сервис закрыт. См. config_local.example.py", 503)
    auth = request.authorization
    ok = (auth is not None
          and hmac.compare_digest(auth.username or "", user)
          and hmac.compare_digest(auth.password or "", password))
    if not ok:
        return Response("Нужна авторизация", 401,
                        {"WWW-Authenticate": 'Basic realm="SEO Meta Auditor"'})
    return None


def _result_path():
    token = session.get("rid", "")
    return os.path.join(RESULTS_DIR, token + ".json") if re.fullmatch(r"[0-9a-f]{32}", token) else None


def save_rows(rows):
    """Сохраняет результат проверки для текущего пользователя (по cookie-сессии)."""
    os.makedirs(RESULTS_DIR, exist_ok=True)
    now = time.time()
    for name in os.listdir(RESULTS_DIR):
        full = os.path.join(RESULTS_DIR, name)
        if name.endswith(".json") and now - os.path.getmtime(full) > RESULTS_TTL:
            os.remove(full)
    if not re.fullmatch(r"[0-9a-f]{32}", session.get("rid", "")):
        session["rid"] = secrets.token_hex(16)
    with open(_result_path(), "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False)


def load_rows():
    path = _result_path()
    if not path or not os.path.exists(path):
        return []
    with open(path, encoding="utf-8") as f:
        return json.load(f)


@app.route("/", methods=["GET", "POST"])
def index():
    rows, error, summary = [], None, None
    if request.method == "POST":
        try:
            sitemap = request.form.get("sitemap", "").strip()
            if sitemap:
                urls = auditor.urls_from_sitemap(sitemap)
            else:
                urls = auditor.parse_urls(request.form.get("urls", ""))
            if not urls:
                error = "Добавьте хотя бы один URL или адрес sitemap.xml"
            else:
                rows = auditor.run_audit(urls)
                save_rows(rows)
                summary = {lvl: sum(r["level"] == lvl for r in rows)
                           for lvl in ("critical", "warning", "info", "ok")}
                summary["total"] = len(rows)
        except auditor.UnsafeURL as exc:
            error = f"Адрес sitemap не разрешён: {exc}"
        except requests.RequestException as exc:
            error = f"Не удалось загрузить sitemap: {type(exc).__name__}"
    return render_template("index.html", rows=rows, error=error, summary=summary,
                           max_urls=auditor.MAX_URLS)


@app.route("/export")
def export():
    rows = load_rows()
    if not rows:
        return "Сначала запустите проверку", 400
    wb = Workbook()
    ws = wb.active
    ws.title = "SEO аудит"
    head = ["URL", "Код", "Title", "Длина title", "Description", "Длина desc",
            "H1", "Кол-во H1", "Canonical", "Проблемы"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r["url"], r["status"], r["title"], len(r["title"]), r["description"],
                   len(r["description"]), r["h1"], r["h1_count"], r["canonical"],
                   "; ".join(msg for _, msg in r["issues"])])
        fill = PatternFill("solid", fgColor=FILLS[r["level"]])
        ws.cell(row=ws.max_row, column=10).fill = fill
    for i, width in enumerate([45, 7, 45, 8, 55, 8, 35, 8, 40, 50], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="seo_audit.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    # локальный запуск: без пароля сервис доступен только с этого компьютера
    if not setting("APP_PASSWORD"):
        app.config["AUTH_REQUIRED"] = False
    app.run(debug=False, host="127.0.0.1", port=5000)
